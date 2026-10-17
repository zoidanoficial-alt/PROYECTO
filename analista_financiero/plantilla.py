"""Plantilla de Excel para que los clientes entreguen su información.

Usa exactamente los nombres de columna que el analista reconoce, así que un
archivo llenado con esta plantilla se lee sin preguntas.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .marca import NOMBRE, TINTA

HOJAS = {
    "Ventas": {
        "columnas": ["Fecha", "Folio", "Cliente", "Producto", "Cantidad", "Importe", "Moneda", "Fecha de cobro"],
        "ejemplo": [["05/04/2026", "F-1001", "Cliente A", "Producto 1", 10, 3200, "MXN", "05/04/2026"],
                    ["12/04/2026", "F-1002", "Cliente B", "Producto 2", 2, 1800, "MXN", ""]],
        "ayuda": "Una fila por venta o factura. «Fecha de cobro» vacía = todavía no te pagan.",
    },
    "Costo de ventas": {
        "columnas": ["Fecha", "Producto", "Cantidad", "Importe", "Fecha de pago"],
        "ejemplo": [["30/04/2026", "Producto 1", 10, 1500, "20/05/2026"],
                    ["30/04/2026", "Producto 2", 2, 1440, "30/04/2026"]],
        "ayuda": "Lo que cuesta producir o comprar lo que vendiste, idealmente por producto.",
    },
    "Gastos": {
        "columnas": ["Fecha", "Concepto", "Categoría", "Importe", "Moneda", "Fecha de pago"],
        "ejemplo": [["01/04/2026", "Renta del local", "Renta", 18000, "MXN", "01/04/2026"],
                    ["15/04/2026", "Nómina quincena", "Nómina", 19000, "MXN", "15/04/2026"]],
        "ayuda": "Gastos de operación: renta, nómina, luz, publicidad, mantenimiento, etc.",
    },
}

INSTRUCCIONES = [
    "Cómo llenar este archivo",
    "",
    "1. Llena las hojas «Ventas», «Costo de ventas» y «Gastos». Borra las filas de ejemplo.",
    "2. Incluye de preferencia los últimos 6 a 12 meses completos.",
    "3. Fechas en formato día/mes/año (05/04/2026). Importes solo con números, sin texto.",
    "4. Si tienes operaciones en dólares u otra moneda, indícalo en la columna «Moneda».",
    "5. «Producto» y «Cantidad» permiten saber qué productos te dejan más dinero.",
    "6. «Fecha de cobro» y «Fecha de pago» permiten calcular el efectivo real (no solo la utilidad).",
    "7. No agregues filas de totales: el análisis los calcula.",
    "",
    "Si tu sistema contable exporta a Excel o CSV, también puedes enviarnos ese archivo tal cual.",
    "Tu información es confidencial y solo se usa para elaborar tu diagnóstico.",
]


def generar_plantilla(ruta: str | Path) -> Path:
    wb = Workbook()
    ins = wb.active
    ins.title = "Instrucciones"
    ins["A1"] = NOMBRE
    ins["A1"].font = Font(size=18, bold=True, color=TINTA[1:])
    for i, linea in enumerate(INSTRUCCIONES, start=3):
        celda = ins.cell(row=i, column=1, value=linea)
        celda.font = Font(bold=(i == 3), size=12 if i == 3 else 11, color=TINTA[1:])
    ins.column_dimensions["A"].width = 110

    encabezado = PatternFill("solid", fgColor=TINTA[1:])
    for nombre, h in HOJAS.items():
        ws = wb.create_sheet(nombre)
        for j, col in enumerate(h["columnas"], start=1):
            c = ws.cell(row=1, column=j, value=col)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = encabezado
            c.alignment = Alignment(horizontal="center")
            ws.column_dimensions[get_column_letter(j)].width = max(14, len(col) + 4)
        for i, fila in enumerate(h["ejemplo"], start=2):
            for j, v in enumerate(fila, start=1):
                ws.cell(row=i, column=j, value=v).font = Font(italic=True, color="7A7A7A")
        # La ayuda va como comentario: una celda junto al encabezado se leería como otra columna.
        ws["A1"].comment = Comment(h["ayuda"], NOMBRE)
        ws.freeze_panes = "A2"
    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta
